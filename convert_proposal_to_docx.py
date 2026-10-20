import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_proposal_docx():
    doc = Document()
    
    # Page Margins: 1 inch everywhere
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Palette
    RED = RGBColor(215, 25, 32)      # Sime Darby Red #D71920
    SLATE = RGBColor(15, 23, 42)     # #0F172A
    TEAL = RGBColor(15, 118, 110)    # Emerald #0F766E
    MUTED = RGBColor(100, 116, 139)  # Slate 500

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("GREENPULSE: THE ZERO-GENSET SITE OPERATING SYSTEM")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RED

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Strategic Innovation Proposal | Sime Darby Property University Hackathon 2026\nTrack: PLANET (Energy Efficiency & Sustainable Living)")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = MUTED
    run_sub.font.bold = True

    # Read markdown file
    with open("PROPOSAL_GREENPULSE_SIME_DARBY.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_data = []

    def flush_table(data):
        if not data:
            return
        # Parse table markdown
        rows = []
        for line in data:
            parts = [c.strip() for c in line.split("|")[1:-1]]
            if parts and not all(c == '-' or c.startswith(':-') or c.startswith('--') for c in parts):
                rows.append(parts)
        
        if not rows:
            return
            
        num_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = table.cell(i, j)
                cell.text = val
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                
                # Format cell
                if i == 0:
                    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
                    cell._tc.get_or_add_tcPr().append(shading)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)
                else:
                    if i % 2 == 1:
                        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                        cell._tc.get_or_add_tcPr().append(shading)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
        
        doc.add_paragraph() # Spacing

    for line in lines:
        raw = line.rstrip()
        
        # Skip title and subtitle handled above
        if raw.startswith("# GREENPULSE") or raw.startswith("## Strategic Innovation") or raw.startswith("### Track:"):
            continue

        if raw.startswith("|"):
            in_table = True
            table_data.append(raw)
            continue
        elif in_table:
            in_table = False
            flush_table(table_data)
            table_data = []

        if raw.startswith("```"):
            continue # skip mermaid/code block fences
            
        if raw.startswith("## "):
            h2 = doc.add_paragraph()
            h2.paragraph_format.space_before = Pt(14)
            h2.paragraph_format.space_after = Pt(4)
            r = h2.add_run(raw.replace("## ", ""))
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RED
            continue
            
        if raw.startswith("### "):
            h3 = doc.add_paragraph()
            h3.paragraph_format.space_before = Pt(10)
            h3.paragraph_format.space_after = Pt(3)
            r = h3.add_run(raw.replace("### ", ""))
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = SLATE
            continue

        if raw.startswith("* ") or raw.startswith("- "):
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(3)
            clean_text = raw[2:]
            # parse bold
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = bp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    bp.add_run(part)
            continue

        if raw.startswith("> "):
            quote = doc.add_paragraph()
            quote.paragraph_format.left_indent = Inches(0.5)
            quote.paragraph_format.space_after = Pt(6)
            quote.paragraph_format.space_before = Pt(6)
            r = quote.add_run(raw.replace("> ", ""))
            r.font.italic = True
            r.font.color.rgb = TEAL
            continue

        if raw.strip() == "---":
            continue

        if raw.strip():
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', raw)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table(table_data)

    target_name = "PROPOSAL_GREENPULSE_SIME_DARBY.docx"
    try:
        doc.save(target_name)
        print(f"{target_name} updated successfully!")
    except PermissionError:
        target_name = "PROPOSAL_GREENPULSE_SIME_DARBY_v2.docx"
        doc.save(target_name)
        print(f"File locked by Word. Saved updated proposal as {target_name} successfully!")

if __name__ == '__main__':
    create_proposal_docx()
